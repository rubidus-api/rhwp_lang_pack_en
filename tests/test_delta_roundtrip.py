"""
변경분 번역 왕복의 불변식을 시험한다 — 진짜 translations/ 는 건드리지 않고 임시 폴더에서.

  python3 -m unittest tests/test_delta_roundtrip.py
"""
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = pathlib.Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import langpack_delta as ld  # noqa: E402

from openpyxl import load_workbook  # noqa: E402  (langpack_delta 가 work/pylib 를 경로에 넣는다)


def run(script, *args):
    # 부모의 환경을 물려준다 — openpyxl 이 어디 깔렸든(사용자 site·PYTHONPATH·work/pylib) 같은 것을 쓴다.
    env = dict(os.environ)
    env['PYTHONPATH'] = os.pathsep.join(p for p in [str(ROOT / 'work' / 'pylib'), env.get('PYTHONPATH', '')] if p)
    return subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)],
                          capture_output=True, text=True, env=env)


class DeltaRoundtrip(unittest.TestCase):
    def setUp(self):
        self.dir = pathlib.Path(tempfile.mkdtemp())
        tr = self.dir / 'translations'
        (tr / 'catalog').mkdir(parents=True)
        self.tr = tr
        ko = {
            'menu.file': '파일',                       # 확정된 번역
            'dialog.a.title': '저장',                  # 확정된 번역, 같은 원문을
            'command.file.save.label': '저장',          # 두 자리가 쓴다
            'dialog.a.text': '{p1}쪽',                  # 검토 안 된 번역
            'dialog.b.title': '새 문구(N)',              # 번역 없음(새로 생김)
            'dialog.b.text': '줄\n바꿈',                 # 번역 없음
        }
        memory = {'파일': 'File', '저장': 'Save', '{p1}쪽': 'page {p1}', '옛 문구': 'Old'}
        reviewed = {'파일': 'File', '저장': 'Save', '옛 문구': 'Old'}
        ld.save_json(tr / 'catalog' / 'ko.json', ko)
        ld.save_json(tr / 'ko-en.json', memory)
        ld.save_json(tr / 'reviewed.json', reviewed)
        self.book = self.dir / 'delta.xlsx'

    def tearDown(self):
        shutil.rmtree(self.dir)

    def export(self):
        result = run('export-delta-xlsx.py', '--translations', self.tr, '--studio', self.dir / 'none',
                     '--out', self.book)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        return load_workbook(self.book)

    def rows(self, wb):
        return {row[ld.C_KO].value: row for row in wb[ld.SHEET].iter_rows(min_row=2)}

    def test_export_contains_only_new_and_unreviewed(self):
        rows = self.rows(self.export())
        self.assertEqual(set(rows), {'{p1}쪽', '새 문구(N)', '줄\n바꿈'})
        self.assertEqual(rows['새 문구(N)'][ld.C_STATE].value, ld.STATE_NEW)
        self.assertEqual(rows['{p1}쪽'][ld.C_STATE].value, ld.STATE_UNREVIEWED)
        self.assertEqual(rows['{p1}쪽'][ld.C_EN].value, 'page {p1}')  # 현재 번역을 미리 채운다

    def test_import_merges_and_never_deletes(self):
        wb = self.export()
        rows = self.rows(wb)
        rows['새 문구(N)'][ld.C_EN].value = 'New text'      # 접근키 (N) 빠짐 → 붙인다
        rows['줄\n바꿈'][ld.C_EN].value = ''                   # 비워 두면 건너뛴다
        # '{p1}쪽' 은 미리 채운 값 그대로 = 확정
        wb.save(self.book)

        result = run('import-delta-xlsx.py', self.book, '--translations', self.tr, '--write')
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        memory = json.loads((self.tr / 'ko-en.json').read_text(encoding='utf-8'))
        reviewed = json.loads((self.tr / 'reviewed.json').read_text(encoding='utf-8'))
        # 시트에 없던 번역(파일·저장·옛 문구)은 그대로 남는다
        self.assertEqual(memory['파일'], 'File')
        self.assertEqual(memory['저장'], 'Save')
        self.assertEqual(memory['옛 문구'], 'Old')
        self.assertEqual(memory['새 문구(N)'], 'New text(N)')
        self.assertNotIn('줄\n바꿈', memory)
        self.assertEqual(reviewed['{p1}쪽'], 'page {p1}')
        self.assertEqual(reviewed['새 문구(N)'], 'New text(N)')

        # 다시 내보내면 비워 둔 것만 남는다
        again = run('export-delta-xlsx.py', '--translations', self.tr, '--check')
        self.assertEqual(again.returncode, 3)
        self.assertIn('변경분 1개', again.stdout)

    def test_placeholder_mismatch_rejects_whole_file(self):
        wb = self.export()
        rows = self.rows(wb)
        rows['{p1}쪽'][ld.C_EN].value = 'page'               # {p1} 이 사라짐
        rows['새 문구(N)'][ld.C_EN].value = 'New(N)'
        wb.save(self.book)
        before = (self.tr / 'ko-en.json').read_bytes()
        result = run('import-delta-xlsx.py', self.book, '--translations', self.tr, '--write')
        self.assertEqual(result.returncode, 2)
        self.assertIn('자리표시자', result.stdout)
        self.assertEqual((self.tr / 'ko-en.json').read_bytes(), before)  # 한 줄도 안 쓴다

    def test_edited_source_cell_is_recovered_by_fingerprint(self):
        wb = self.export()
        rows = self.rows(wb)
        row = rows['새 문구(N)']
        row[ld.C_KO].value = '새 문구 (고쳐 버림)'
        row[ld.C_EN].value = 'New(N)'
        wb.save(self.book)
        result = run('import-delta-xlsx.py', self.book, '--translations', self.tr, '--write')
        self.assertEqual(result.returncode, 0, result.stdout)
        memory = json.loads((self.tr / 'ko-en.json').read_text(encoding='utf-8'))
        self.assertEqual(memory['새 문구(N)'], 'New(N)')
        self.assertNotIn('새 문구 (고쳐 버림)', memory)

    def test_excel_autocorrect_is_reverted(self):
        self.assertEqual(ld.fix_autocorrect('글자처럼 취급(C)', 'Treat as character©'), 'Treat as character(C)')
        self.assertEqual(ld.fix_autocorrect('열기...', 'Open…'), 'Open...')

    def test_newline_survives_excel_roundtrip(self):
        wb = self.export()
        self.rows(wb)['줄\n바꿈'][ld.C_EN].value = 'Line\r\nbreak'
        wb.save(self.book)
        result = run('import-delta-xlsx.py', self.book, '--translations', self.tr, '--write')
        self.assertEqual(result.returncode, 0, result.stdout)
        memory = json.loads((self.tr / 'ko-en.json').read_text(encoding='utf-8'))
        self.assertEqual(memory['줄\n바꿈'], 'Line\nbreak')


if __name__ == '__main__':
    unittest.main()
