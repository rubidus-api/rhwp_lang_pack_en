#!/usr/bin/env python3
"""
번역해 온 변경분 엑셀을 번역 정본에 **병합**한다.

  python3 scripts/import-delta-xlsx.py review/delta-….xlsx           # 무엇이 바뀔지 보고만
  python3 scripts/import-delta-xlsx.py review/delta-….xlsx --write   # translations/ 에 기록

코드에 반영하려면 그다음 `sh scripts/regenerate.sh` — 한 번에 하려면 `sh scripts/apply-translations.sh`.

규칙
  - 시트에 없는 원문은 건드리지 않는다(병합만). 번역 메모리를 새로 짜지 않는다.
  - 번역 칸이 빈 행은 건너뛴다 — 다음 변경분에 다시 나온다.
  - 원문 칸이 고쳐졌으면 숨긴 지문 열로 원래 원문을 되찾는다. 못 찾으면 그 행은 거부한다.
  - 엑셀 자동고침((C)→© 등)은 되돌리고, 원문 끝의 접근키 (X) 가 빠졌으면 붙인다.
  - 자리표시자({p1} 등)가 원문과 다르면 **그 행을 거부**한다. 화면에서 값이 사라지기 때문이다.
    거부가 하나라도 있으면 --write 해도 아무것도 쓰지 않는다(종료코드 2) — 반쯤 들어간 상태를 만들지 않는다.
  - 반영한 행은 translations/reviewed.json 에 '사람이 확정한 번역' 으로 적는다.
"""
import argparse
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import langpack_delta as ld  # noqa: E402


def main(argv):
    parser = argparse.ArgumentParser(description='번역해 온 변경분 엑셀을 번역 정본에 병합한다')
    parser.add_argument('workbook', type=pathlib.Path)
    parser.add_argument('--write', action='store_true', help='translations/ 에 기록한다')
    parser.add_argument('--translations', type=pathlib.Path, default=ld.TRANSLATIONS)
    opts = parser.parse_args(argv)
    write, translations, workbook_path = opts.write, opts.translations, opts.workbook

    try:
        from openpyxl import load_workbook
    except ImportError:
        print('import-delta: openpyxl 이 필요하다 — pip install --user openpyxl', file=sys.stderr)
        return 1

    wb = load_workbook(workbook_path)
    if ld.SHEET not in wb.sheetnames:
        print(f"import-delta: '{ld.SHEET}' 시트가 없다 — export-delta-xlsx.py 가 만든 파일인가?", file=sys.stderr)
        return 1
    ws = wb[ld.SHEET]
    header = [ld.normalize_cell(c.value) for c in ws[1]]
    if header[:len(ld.HEADERS)] != ld.HEADERS:
        print('import-delta: 열 머리가 다르다 — 열을 옮기거나 지우지 말 것', file=sys.stderr)
        print(f'  기대: {ld.HEADERS}\n  실제: {header}', file=sys.stderr)
        return 1

    memory_path = translations / 'ko-en.json'
    reviewed_path = translations / 'reviewed.json'
    memory = ld.load_json(memory_path)
    reviewed = ld.load_json(reviewed_path)
    ko = ld.load_json(translations / 'catalog' / 'ko.json')
    by_fingerprint = {ld.fingerprint(v): v for v in set(ko.values()) | set(memory)}

    applied, unchanged, confirmed, skipped, rejected, notes = [], [], [], [], [], []
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue
        cells = list(row) + [None] * (len(ld.HEADERS) - len(row))
        korean = ld.normalize_cell(cells[ld.C_KO])
        fp = ld.normalize_cell(cells[ld.C_FP]).strip()
        english = ld.normalize_cell(cells[ld.C_EN])

        if fp:
            original = by_fingerprint.get(fp)
            if original is None:
                rejected.append((number, korean, '지문에 맞는 원문이 없다 — 오래된 파일이거나 행이 섞였다'))
                continue
            if original != korean:
                notes.append((number, f'원문 칸이 고쳐져 있어 원래 원문으로 되돌림: {original!r}'))
                korean = original
        elif not korean:
            continue

        if english.strip() == '':
            skipped.append((number, korean))
            continue

        fixed = ld.fix_autocorrect(korean, english)
        if fixed != english:
            notes.append((number, f'엑셀 자동고침 되돌림: {english!r} → {fixed!r}'))
        with_key = ld.keep_access_key(korean, fixed)
        if with_key != fixed:
            notes.append((number, f'접근키 표시 붙임: {fixed!r} → {with_key!r}'))
        english = with_key

        errors, warnings = ld.check_translation(korean, english)
        if errors:
            rejected.append((number, korean, '; '.join(errors)))
            continue
        for warning in warnings:
            notes.append((number, f'{korean!r}: {warning}'))
        if korean not in set(ko.values()):
            notes.append((number, f'{korean!r}: 지금 코드가 쓰지 않는 원문이다(상류에서 사라짐?) — 메모리에만 남긴다'))

        previous = memory.get(korean)
        if previous == english:
            (unchanged if reviewed.get(korean) == english else confirmed).append((number, korean, english))
        else:
            applied.append((number, korean, previous, english))

    print(f'{workbook_path.name}: 새로·고친 번역 {len(applied)}, 그대로 확정 {len(confirmed)}, '
          f'이미 확정 {len(unchanged)}, 빈칸 건너뜀 {len(skipped)}, 거부 {len(rejected)}')
    for number, korean, previous, english in applied[:60]:
        before = '(새 문구)' if previous is None else repr(previous)
        print(f'  {number}행 {korean!r}: {before} → {english!r}')
    if len(applied) > 60:
        print(f'  … 외 {len(applied) - 60}건')
    for number, message in notes:
        print(f'  참고 {number}행: {message}')
    for number, korean, reason in rejected:
        print(f'  ★거부 {number}행 {korean!r}: {reason}')

    if rejected:
        print('import-delta: 거부된 행이 있어 아무것도 쓰지 않았다 — 엑셀을 고쳐 다시 실행할 것', file=sys.stderr)
        return 2
    if not write:
        print('import-delta: 보고만 했다. 기록하려면 --write')
        return 0

    before_count = len(memory)
    for _, korean, _, english in applied:
        memory[korean] = english
    for _, korean, *rest in applied + confirmed + unchanged:
        reviewed[korean] = memory[korean]
    # 병합 불변식: 메모리는 늘거나 그대로일 뿐 줄지 않는다.
    assert len(memory) >= before_count, '번역 메모리가 줄었다 — 병합 규칙 위반'
    ld.save_json(memory_path, memory)
    ld.save_json(reviewed_path, reviewed)
    print(f'import-delta: 기록함 — 번역 메모리 {before_count} → {len(memory)}, 확정 {len(reviewed)}')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
