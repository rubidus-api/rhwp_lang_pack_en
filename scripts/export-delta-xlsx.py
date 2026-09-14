#!/usr/bin/env python3
"""
번역할 변경분만 엑셀로 모은다.

  python3 scripts/export-delta-xlsx.py                 # review/delta-YYYYMMDD-HHMM.xlsx
  python3 scripts/export-delta-xlsx.py --out 파일.xlsx
  python3 scripts/export-delta-xlsx.py --check          # 파일은 안 만들고 개수만 (변경분 있으면 종료코드 3)

한 행은 한 **원문 값**이다. 같은 원문을 여러 자리가 쓰면 키 칸에 모두 적는다 — 번역 메모리가
값 단위이기 때문이다. 자리마다 달라야 하는 번역은 translations/ko-en-overrides.json 에 키로 적는다.

'번역(영어)' 칸은 지금 쓰이는 번역(없으면 빈칸)으로 미리 채운다. 고칠 것만 고치고, 맞으면
그대로 두면 된다 — 되돌려 받은 파일의 채워진 행은 전부 '사람이 확정한 번역' 으로 기록된다.
확정하지 않을 행은 번역 칸을 비우면 다음 변경분에 다시 나온다.
"""
import argparse
import datetime
import importlib.util
import pathlib
import subprocess
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import langpack_delta as ld  # noqa: E402


def load_review_helpers():
    """화면 위치 계산은 전체 검토표 도구와 같은 규칙을 쓴다(파일명에 '-' 가 있어 importlib 로)."""
    path = ld.ROOT / 'scripts' / 'export-review-xlsx.py'
    spec = importlib.util.spec_from_file_location('export_review_xlsx', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def upstream_revision(repo):
    try:
        return subprocess.run(['git', '-C', str(repo), 'rev-parse', '--short', 'HEAD'],
                              capture_output=True, text=True, check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return '(알 수 없음)'


def main(argv):
    parser = argparse.ArgumentParser(description='번역할 변경분만 엑셀로 모은다')
    parser.add_argument('--out', type=pathlib.Path, help='기본: review/delta-YYYYMMDD-HHMM.xlsx')
    parser.add_argument('--check', action='store_true', help='개수만 센다. 변경분이 있으면 종료코드 3')
    parser.add_argument('--translations', type=pathlib.Path, default=ld.TRANSLATIONS)
    parser.add_argument('--studio', type=pathlib.Path, default=ld.ROOT / 'work/rhwp/rhwp-studio')
    opts = parser.parse_args(argv)
    check_only, out_path, translations, studio = opts.check, opts.out, opts.translations, opts.studio

    ko = ld.load_json(translations / 'catalog' / 'ko.json')
    memory = ld.load_json(translations / 'ko-en.json')
    reviewed = ld.load_json(translations / 'reviewed.json')
    keep = ld.load_json(translations / 'keep-as-is.json')
    if not ko:
        print('export-delta: 카탈로그가 비었다 — 먼저 sh scripts/regenerate.sh', file=sys.stderr)
        return 1

    delta = ld.delta_values(ko, memory, reviewed)
    new = sum(1 for s in delta.values() if s == ld.STATE_NEW)
    print(f'변경분 {len(delta)}개 (새 문구 {new}, 검토 필요 {len(delta) - new}) / 쓰이는 원문 {len(set(ko.values()))}개')
    if check_only:
        return 3 if delta else 0
    if not delta:
        print('export-delta: 번역할 것이 없다')
        return 0

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print('export-delta: openpyxl 이 필요하다 — pip install --user openpyxl '
              '(또는 pip install --target work/pylib openpyxl)', file=sys.stderr)
        return 1

    helpers = load_review_helpers()
    keys_by_value = {}
    for key, value in ko.items():
        keys_by_value.setdefault(value, []).append(key)
    uses = helpers.find_uses(studio, set(ko)) if studio.is_dir() else {}
    baseline = ld.ROOT / 'work' / 'ui.baseline'
    titles = helpers.dialog_titles(ko, baseline)
    keep_values, keep_keys = keep.get('values', {}), keep.get('keys', {})

    rows = []
    for korean, state in delta.items():
        keys = sorted(keys_by_value[korean])
        places = []
        files = set()
        for key in keys:
            place, role, used_in = helpers.where(key, ko, titles, uses)
            label = f'{place} {role}'.strip()
            if label not in places:
                places.append(label)
            files.update(f for f in used_in.split(', ') if f and f != '—')
        notes = []
        if korean in keep_values:
            notes.append('한국어로 둠 — ' + keep_values[korean])
        if any(k in keep_keys for k in keys):
            notes.append('일부 자리는 한국어로 둠(keep-as-is.json)')
        if ld.PLACEHOLDER.search(korean):
            notes.append('{p1} 같은 자리표시자는 코드가 채우는 값 — 번역에도 그대로 둘 것')
        if '\n' in korean:
            notes.append('줄바꿈 수를 원문과 맞출 것(Alt+Enter)')
        if ld.ACCESS_KEY.search(korean):
            notes.append('끝의 (X) 접근키 표시는 남길 것 — 빠뜨리면 반영 도구가 붙인다')
        order = min(helpers.AREA_ORDER.get(k.split('.')[0], 9) for k in keys)
        current = memory.get(korean, '')
        rows.append((0 if state == ld.STATE_NEW else 1, order, places[0] if places else '', korean, state,
                     '\n'.join(places[:4]) + (f'\n… 외 {len(places) - 4}곳' if len(places) > 4 else ''),
                     current, '\n'.join(keys), ', '.join(sorted(files)), '\n'.join(notes)))
    rows.sort(key=lambda r: r[:3])

    wb = Workbook()
    ws = wb.active
    ws.title = ld.SHEET
    ws.append(ld.HEADERS)
    for index, (_, _, _, korean, state, places, current, keys, files, notes) in enumerate(rows, start=1):
        ws.append([index, state, places, korean, current, current, keys, files, notes, ld.fingerprint(korean)])

    head_fill = PatternFill('solid', fgColor='DDE5F0')
    input_fill = PatternFill('solid', fgColor='FFF7D6')
    new_fill = PatternFill('solid', fgColor='FDE2E1')
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    widths = [6, 10, 30, 44, 44, 30, 36, 24, 34, 14]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        row[ld.C_EN].fill = input_fill
        if row[ld.C_STATE].value == ld.STATE_NEW:
            row[ld.C_STATE].fill = new_fill
        # 엑셀이 '=…' 로 시작하는 문구를 수식으로 읽지 않게 문자열로 못 박는다
        for cell in (row[ld.C_KO], row[ld.C_EN], row[ld.C_CURRENT]):
            cell.data_type = 's'
    ws.column_dimensions[get_column_letter(ld.C_FP + 1)].hidden = True
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(ld.HEADERS))}{ws.max_row}'

    info = wb.create_sheet(ld.INFO_SHEET)
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
    for line in [
        ['항목', '값'],
        ['만든 때', now],
        ['상류 기준', upstream_revision(studio.parent)],
        ['변경분', f'{len(rows)}개 (새 문구 {new}, 검토 필요 {len(rows) - new})'],
        ['하는 법', "노란 '번역(영어)' 칸만 고친다. 맞으면 그대로 둔다. 확정하지 않을 행은 번역 칸을 비운다."],
        ['되돌리는 법', 'sh scripts/apply-translations.sh <이 파일>  — 병합·검사 후 코드까지 다시 만든다'],
        ['건드리지 말 것', "'원문(한국어)' 칸과 숨긴 '원문 지문' 열. 시트 이름도 그대로 둔다."],
    ]:
        info.append(line)
    for cell in info[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    info.column_dimensions['A'].width = 16
    info.column_dimensions['B'].width = 90

    if out_path is None:
        stamp = datetime.datetime.now().strftime('%Y%m%d-%H%M')
        out_path = ld.ROOT / 'review' / f'delta-{stamp}.xlsx'
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    try:
        shown = out_path.resolve().relative_to(ld.ROOT)
    except ValueError:
        shown = out_path.name
    print(f'export-delta: {shown} ({len(rows)}행)')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
